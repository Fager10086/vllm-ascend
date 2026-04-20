#!/usr/bin/env python
# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt
import sys

print("Starting...", file=sys.stderr)

prs = Presentation()
print("Created prs", file=sys.stderr)

prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
print("Set dimensions", file=sys.stderr)

slide = prs.slides.add_slide(prs.slide_layouts[6])
print("Added slide", file=sys.stderr)

tb = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(1))
p = tb.text_frame.paragraphs[0]
p.text = "Test Slide"
p.font.size = Pt(36)
print("Added text", file=sys.stderr)

output = "D:/vllm-ascend/test.pptx"
print(f"Saving to {output}...", file=sys.stderr)

try:
    prs.save(output)
    print("Save called", file=sys.stderr)
except Exception as e:
    print(f"Error: {e}", file=sys.stderr)
    import traceback

    traceback.print_exc(file=sys.stderr)

import os

print(f"File exists: {os.path.exists(output)}", file=sys.stderr)

if os.path.exists(output):
    print(f"File size: {os.path.getsize(output)}", file=sys.stderr)
